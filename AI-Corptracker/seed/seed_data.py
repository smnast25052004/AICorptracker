"""
Seed script — populates the database from 'Примеры целей.xlsx'.

Loads all goals, key results (as projects+tasks), and employees
from the real corporate dataset.
"""
import random
from datetime import datetime
from pathlib import Path

import openpyxl

from shared.database import engine, get_db_session
from shared.models import Base
from shared.models.employee import Employee
from shared.models.goal import StrategicGoal, GoalStatus, GoalPriority
from shared.models.project import Project, ProjectStatus
from shared.models.task import Task, TaskStatus, TaskPriority
from shared.models.document import Document, DocumentStatus, DocumentType
from processor.ai.embeddings import generate_embedding


XLSX_PATH = Path(__file__).resolve().parent.parent / "Примеры целей.xlsx"


def _parse_xlsx():
    """Parse the XLSX file and return structured data."""
    wb = openpyxl.load_workbook(str(XLSX_PATH))
    ws = wb.active

    employees_map = {}
    goals_map = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        (emp_code, fio, post, dept, div, gid, gname, kr_id, kr_name,
         gt_code, gt, period_end, period_type, period_year, period_code,
         goal_weight, prozt, prozt_fact, result_plan, result_fact,
         tracking_sys, tracking_link, description, summary, board,
         key_product, entity) = row

        if fio and fio.strip() and emp_code:
            fio_clean = fio.strip()
            if fio_clean not in employees_map:
                employees_map[fio_clean] = {
                    "emp_code": emp_code,
                    "post": post,
                    "department": dept,
                    "division": div,
                }

        if not gid:
            continue

        if gid not in goals_map:
            progress_fact = prozt_fact if prozt_fact is not None else 0
            if isinstance(progress_fact, (int, float)):
                progress_val = min(float(progress_fact), 100.0)
            else:
                progress_val = 0.0

            goals_map[gid] = {
                "name": gname,
                "type": gt,
                "board": board,
                "period_end": str(period_end)[:10] if period_end else "2025-12-31",
                "progress": progress_val,
                "owner": None,
                "key_results": {},
            }

        if fio:
            goals_map[gid]["owner"] = fio.strip()

        if kr_id and kr_name:
            if kr_id not in goals_map[gid]["key_results"]:
                goals_map[gid]["key_results"][kr_id] = {
                    "name": kr_name,
                    "result_plan": result_plan,
                    "result_fact": result_fact,
                    "tracking": tracking_sys,
                }

    return employees_map, goals_map


def _goal_type_to_priority(goal_type):
    """Map goal_type from XLSX to GoalPriority."""
    if goal_type in ("КПЭ", "Метрика"):
        return GoalPriority.HIGH
    return GoalPriority.MEDIUM


def _progress_to_status(progress):
    """Derive goal status from progress percentage."""
    if progress >= 100:
        return GoalStatus.COMPLETED
    if progress >= 80:
        return GoalStatus.ON_TRACK
    if progress >= 50:
        return GoalStatus.AT_RISK
    return GoalStatus.CRITICAL


def _kr_to_task_status(result_plan, result_fact):
    """Derive task status from key result plan/fact values."""
    if result_plan is None and result_fact is None:
        return TaskStatus.IN_PROGRESS
    if result_fact is not None and result_plan is not None:
        try:
            if float(result_fact) >= float(result_plan):
                return TaskStatus.DONE
            if float(result_fact) >= float(result_plan) * 0.8:
                return TaskStatus.IN_REVIEW
            return TaskStatus.IN_PROGRESS
        except (ValueError, TypeError):
            return TaskStatus.IN_PROGRESS
    if result_fact is not None:
        return TaskStatus.DONE
    return TaskStatus.TODO


def seed():
    print("Creating database tables...")
    Base.metadata.create_all(bind=engine)

    with get_db_session() as db:
        existing = db.query(StrategicGoal).count()
        if existing > 0:
            print(f"Database already has {existing} goals. Skipping seed.")
            return

        if not XLSX_PATH.exists():
            print(f"ERROR: File not found: {XLSX_PATH}")
            return

        print(f"Loading data from {XLSX_PATH.name}...")
        employees_map, goals_map = _parse_xlsx()

        print("Seeding employees...")
        emp_objects = {}
        for i, (fio, info) in enumerate(employees_map.items(), 1):
            email = f"user{info['emp_code']}@company.com"
            emp = Employee(
                full_name=fio,
                email=email,
                department=info["department"] or "Корпоративно-инвестиционный бизнес",
                position=info["post"] or "Специалист",
            )
            db.add(emp)
            emp_objects[fio] = emp
        db.flush()

        print("Seeding strategic goals...")
        goal_objects = {}
        for gid, gdata in goals_map.items():
            goal = StrategicGoal(
                title=gdata["name"],
                description=f"Тип: {gdata['type']}. Board: {gdata['board']}.",
                owner=gdata["owner"] or "Руководство",
                priority=_goal_type_to_priority(gdata["type"]),
                status=_progress_to_status(gdata["progress"]),
                progress=gdata["progress"],
                target_date=gdata["period_end"],
            )
            db.add(goal)
            goal_objects[gid] = goal
        db.flush()

        print("Seeding projects and tasks from key results...")
        project_count = 0
        task_count = 0

        for gid, gdata in goals_map.items():
            goal = goal_objects[gid]
            krs = gdata["key_results"]
            if not krs:
                continue

            proj = Project(
                title=gdata["name"],
                status=ProjectStatus.COMPLETED if gdata["progress"] >= 100 else ProjectStatus.ACTIVE,
                progress=gdata["progress"],
                lead=gdata["owner"] or "Руководство",
                goal_id=goal.id,
            )
            db.add(proj)
            db.flush()
            project_count += 1

            for kr_id, kr_data in krs.items():
                assignee = None
                owner_name = gdata["owner"]
                if owner_name and owner_name in emp_objects:
                    assignee = emp_objects[owner_name]

                task_status = _kr_to_task_status(kr_data["result_plan"], kr_data["result_fact"])

                desc_parts = []
                if kr_data["result_plan"] is not None:
                    desc_parts.append(f"План: {kr_data['result_plan']}")
                if kr_data["result_fact"] is not None:
                    desc_parts.append(f"Факт: {kr_data['result_fact']}")
                if kr_data["tracking"]:
                    desc_parts.append(f"Tracking: {kr_data['tracking']}")
                description = "; ".join(desc_parts) if desc_parts else f"Key Result: {kr_data['name']}"

                task = Task(
                    title=kr_data["name"],
                    status=task_status,
                    priority=TaskPriority.HIGH if gdata["type"] == "КПЭ" else TaskPriority.MEDIUM,
                    source_system="jira",
                    external_id=f"KR-{kr_id}",
                    story_points=random.choice([2, 3, 5, 8]),
                    assignee_id=assignee.id if assignee else None,
                    project_id=proj.id,
                    description=description,
                )
                db.add(task)
                task_count += 1

        db.flush()

        print("Seeding documents...")
        doc_configs = [
            ("Стратегия корпоративно-инвестиционного бизнеса", "report", "approved", "confluence",
             "Стратегический документ описывает ключевые цели и направления развития КИБ."),
            ("Методика расчёта KPI подразделений", "specification", "approved", "confluence",
             "Методика описывает расчёт ключевых показателей эффективности для всех подразделений блока."),
            ("Отчёт по AI-трансформации Q4 2025", "report", "review", "confluence",
             "Результаты внедрения AI-агентов: обучение, аналитика, автоматизация рутинных операций."),
            ("Регламент управления рисками КИБ", "specification", "approved", "edo",
             "Регламент описывает процессы управления кредитными, рыночными и операционными рисками."),
            ("Контракт на облачную инфраструктуру", "contract", "review", "edo",
             "Договор на предоставление облачных вычислительных ресурсов для аналитических систем."),
            ("ТЗ на инвестиционный портал", "specification", "approved", "confluence",
             "Техническое задание на разработку портала анализа данных с AI-функциями."),
            ("Отчёт по удовлетворённости клиентов CSI", "report", "approved", "confluence",
             "Результаты опроса клиентов по качеству обслуживания: РКО, депозиты, бизнес-карты."),
            ("Аналитика международных рынков", "report", "review", "confluence",
             "Исследование потенциала международных расчётов и возможностей на рынках Индии."),
            ("Политика информационной безопасности", "specification", "approved", "edo",
             "Политика кибербезопасности: стандарты, контроль доступа, защита данных."),
            ("SLA информационных систем", "contract", "approved", "edo",
             "Соглашение об уровне обслуживания: доступность 99%, время отклика, надёжность транзакций."),
        ]

        emp_list = list(emp_objects.values())
        for title, doc_type, status, source, content in doc_configs:
            embedding = generate_embedding(content)
            doc = Document(
                title=title,
                content=content,
                doc_type=DocumentType(doc_type),
                status=DocumentStatus(status),
                source_system=source,
                author=random.choice(emp_list).full_name,
                embedding=embedding,
            )
            db.add(doc)

        db.flush()
        print("Seed complete!")
        print(f"  Goals: {len(goal_objects)}")
        print(f"  Projects: {project_count}")
        print(f"  Tasks (key results): {task_count}")
        print(f"  Employees: {len(emp_objects)}")
        print(f"  Documents: {len(doc_configs)}")


if __name__ == "__main__":
    seed()
